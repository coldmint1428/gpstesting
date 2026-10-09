"""Geometry audit for the ER diagram deck: catches overlaps and out-of-bounds shapes.

Run:  <python> docs/audit_er_pptx.py
"""
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0
prs = Presentation("docs/er-diagram.pptx")
SW = prs.slide_width / EMU_IN
SH = prs.slide_height / EMU_IN
print(f"slide size {SW:.2f} x {SH:.2f} in")


def box(s):
    if s.left is None or s.top is None:
        return None
    w = s.width or 0
    h = s.height or 0
    return (s.left / EMU_IN, s.top / EMU_IN,
            (s.left + w) / EMU_IN, (s.top + h) / EMU_IN)


def overlap(a, b, tol=0.005):
    return (min(a[2], b[2]) - max(a[0], b[0]) > tol and
            min(a[3], b[3]) - max(a[1], b[1]) > tol)


problem = 0
for idx, slide in enumerate(prs.slides, start=1):
    print(f"\n--- slide {idx} ---")
    items = []
    for s in slide.shapes:
        b = box(s)
        if b is None:
            continue
        label = (s.text_frame.text.replace("\n", " ")[:40]
                 if s.has_text_frame and s.text_frame.text else "")
        items.append((s.shape_type, s.name, label, b, s.has_text_frame))

    # out of bounds
    for st, name, label, b, _ in items:
        if b[0] < -0.001 or b[1] < -0.001 or b[2] > SW + 0.001 or b[3] > SH + 0.001:
            print(f"  OUT OF BOUNDS {name} '{label}' -> "
                  f"({b[0]:.2f},{b[1]:.2f})-({b[2]:.2f},{b[3]:.2f})")
            problem += 1

    # text boxes that collide with another text-bearing box
    texts = [it for it in items if it[4] and it[2].strip()]
    for i in range(len(texts)):
        for j in range(i + 1, len(texts)):
            if overlap(texts[i][3], texts[j][3]):
                print(f"  TEXT OVERLAP: '{texts[i][2]}'  <->  '{texts[j][2]}'")
                problem += 1

print(f"\n{problem} geometry problem(s)")
